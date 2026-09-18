"""Регрессии на PostgreSQL и HTTP; запускаются в стеке scripts/compose.test.yml."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import HTTPRedirectHandler, Request, build_opener

from openpyxl import load_workbook
from sqlalchemy import delete, event, func, select, text

from app.captcha.engine import KINDS, consume_challenge, issue_challenge, verify_challenge
from app.config import settings
from app.db import SessionLocal, engine
from app.models import CaptchaChallenge, EntryStatus, Purpose, QueueEntry, QueueSession
from app.security import make_admin_token, name_key, new_entry_token, write_entry_token
from app.services import export, queue
from app.templating import templates


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):
        return None


def http(path, data=None, cookie=""):
    request = Request(
        "http://127.0.0.1:8000" + path,
        data=urlencode(data).encode() if data is not None else None,
        headers={"Cookie": cookie} if cookie else {},
    )
    try:
        response = build_opener(NoRedirect).open(request, timeout=10)
    except HTTPError as exc:
        response = exc
    with response:
        return response.code, response.headers, response.read()


class RegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if engine.url.database != "au_queue_test" or engine.url.password != "test-only-password":
            raise RuntimeError("Запускайте регрессии через scripts/smoke.sh в тестовой БД")

    def setUp(self):
        self.sessions = []
        self.captchas = []
        self.admin = "au_admin=" + make_admin_token()
        with SessionLocal() as db:
            self.purpose_id = db.scalar(
                select(Purpose.id).where(Purpose.needs_comment.is_(False))
            )

    def tearDown(self):
        # Только собственные фикстуры: история smoke-сценария остаётся нетронутой.
        with SessionLocal() as db:
            db.execute(delete(QueueSession).where(QueueSession.id.in_(self.sessions)))
            db.execute(delete(CaptchaChallenge).where(CaptchaChallenge.id.in_(self.captchas)))
            db.commit()

    def open(self):
        with SessionLocal() as db:
            session = queue.open_session(db, room="1215", time_from=None, time_to=None, note="Тест")
            self.sessions.append(session.id)
            return session

    def close(self, session):
        with SessionLocal() as db:
            return queue.close_session(db, db.get(QueueSession, session.id))

    def set_window(self, session, *, opened_hours_ago=2, ends_in_minutes=-60):
        """Сдвинуть приём в прошлое: и открытие, и заявленное «до».

        Двигаем именно пару, а не одно «до»: иначе около полуночи приём
        оказался бы «через полночь» и время бы ещё не вышло.
        """
        now = queue.now_local()
        opened = now - dt.timedelta(hours=opened_hours_ago)
        ends = now + dt.timedelta(minutes=ends_in_minutes)
        with SessionLocal() as db:
            stored = db.get(QueueSession, session.id)
            stored.opened_at = opened.astimezone(dt.timezone.utc)
            stored.time_from = opened.time()
            stored.time_to = ends.time()
            db.commit()
            self.assertTrue(queue.joining_closed(stored) == (ends_in_minutes < 0))

    def captcha(self, solved=True):
        with SessionLocal() as db:
            challenge = issue_challenge(db, "quiz")
            self.captchas.append(challenge.id)
            if solved:
                self.assertTrue(verify_challenge(db, challenge.id, str(challenge.answer["index"]))["ok"])
            return challenge

    @staticmethod
    def solution(challenge) -> str:
        """Ответ задания так, как его отправил бы браузер."""
        answer = challenge.answer
        if "values" in answer:
            values = answer["values"]
            if challenge.kind == "truth_myth":
                return ",".join("1" if value else "0" for value in values)
            return ",".join(str(value) for value in values)
        if "node" in answer:
            return answer["node"]
        if "prefix" in answer:
            return str(answer["prefix"])
        return str(answer["index"])

    def add(self, db, session, name="Тестов Иван", comment=""):
        return queue.join_queue(
            db, session=session, full_name=name, full_name_key=name_key(name),
            group_name="КТ-24-04", purpose_id=self.purpose_id, comment=comment,
            token=new_entry_token(),
        )

    def form(self, session, challenge, name="Тестов Иван", comment=""):
        return {"session_id": session.id, "captcha_id": challenge.id, "full_name": name,
                "group_name": "КТ-24-04", "purpose_id": self.purpose_id, "comment": comment}

    def count(self, session):
        with SessionLocal() as db:
            return db.scalar(select(func.count(QueueEntry.id)).where(QueueEntry.session_id == session.id))

    def test_captcha_has_only_one_concurrent_consumer(self):
        challenge = self.captcha()
        barrier = threading.Barrier(2)

        def consume(_):
            with SessionLocal() as db:
                # Оба запроса уже держат старое состояние в identity map.
                cached = db.get(CaptchaChallenge, challenge.id)
                self.assertFalse(cached.consumed)
                barrier.wait(timeout=5)
                return consume_challenge(db, challenge.id)

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(consume, range(2)))
        self.assertEqual(sorted(results), [False, True])

    def test_parallel_http_joins_cannot_reuse_captcha(self):
        session = self.open()
        challenge = self.captcha()
        barrier = threading.Barrier(2)

        def submit(name):
            barrier.wait(timeout=5)
            return http("/join", self.form(session, challenge, name))[0]

        with ThreadPoolExecutor(max_workers=2) as pool:
            codes = list(pool.map(submit, ["Тестов Иван", "Тестова Анна"]))
        self.assertEqual(sorted(codes), [303, 422])
        self.assertEqual(self.count(session), 1)

    def test_parallel_wrong_answers_burn_the_challenge_once(self):
        challenge = self.captcha(solved=False)
        workers = 4
        barrier = threading.Barrier(workers)

        def answer(_):
            with SessionLocal() as db:
                cached = db.get(CaptchaChallenge, challenge.id)
                self.assertEqual(cached.attempts, 0)
                barrier.wait(timeout=5)
                return verify_challenge(db, challenge.id, "-1")

        with ThreadPoolExecutor(max_workers=workers) as pool:
            results = list(pool.map(answer, range(workers)))
        self.assertTrue(all(not result["ok"] for result in results))
        with SessionLocal() as db:
            actual = db.get(CaptchaChallenge, challenge.id)
            # Попытка одна: первый ответ гасит задание, остальные видят погашенное.
            self.assertTrue(actual.consumed)
            self.assertEqual(actual.attempts, 1)
            self.assertTrue(verify_challenge(db, challenge.id, str(actual.answer["index"]))["expired"])

    def test_wrong_answer_is_replaced_with_another_kind(self):
        challenge = self.captcha(solved=False)
        wrong = str((challenge.answer["index"] + 1) % 4)
        code, _, body = http(f"/api/captcha/{challenge.id}", {"answer": wrong})
        result = json.loads(body)
        self.assertEqual(code, 200)
        self.assertFalse(result["ok"])
        self.assertTrue(result["expired"])
        # Разбор ошибки приходит вместе с готовой заменой другого типа.
        self.assertGreater(len(result["text"]), 40)
        replacement = result["replacement"]
        self.captchas.append(replacement["id"])
        self.assertNotEqual(replacement["kind"], challenge.kind)
        self.assertNotIn('"answer"', replacement["html"])
        self.assertIn(f'data-id="{replacement["id"]}"', replacement["html"])

    def test_new_challenge_differs_from_the_one_on_screen(self):
        challenge = self.captcha(solved=False)
        for _ in range(12):
            _, _, body = http(f"/api/captcha?previous={challenge.id}")
            fresh = json.loads(body)
            self.captchas.append(fresh["id"])
            self.assertNotEqual(fresh["kind"], challenge.kind)

    def test_every_kind_builds_and_checks_its_own_answer(self):
        with SessionLocal() as db:
            for kind in KINDS:
                challenge = issue_challenge(db, kind)
                self.captchas.append(challenge.id)
                self.assertEqual(challenge.kind, kind)
                html = templates.get_template("partials/captcha.html").render(
                    request=None, challenge=challenge
                )
                # В разметку уезжает только то, что нужно нарисовать.
                self.assertNotIn('"answer"', html)
                self.assertTrue(
                    verify_challenge(db, challenge.id, self.solution(challenge))["ok"],
                    f"{kind}: собственный ответ не принят",
                )

    def test_join_rechecks_closed_session_with_stale_orm_object(self):
        session = self.open()
        with SessionLocal() as student:
            stale = student.get(QueueSession, session.id)
            self.assertTrue(stale.is_open)
            self.close(session)
            with self.assertRaises(queue.QueueError):
                self.add(student, stale)
        self.assertEqual(self.count(session), 0)

    def test_close_waits_for_join_and_counts_its_entry(self):
        session = self.open()
        ready = threading.Event()
        release = threading.Event()
        closing = threading.Event()
        pid = []

        def hold_commit(db):
            ready.set()
            if not release.wait(timeout=10):
                raise RuntimeError("Тест не освободил транзакцию записи")

        def join():
            with SessionLocal() as db:
                event.listen(db, "before_commit", hold_commit, once=True)
                return self.add(db, db.get(QueueSession, session.id)).id

        def close():
            with SessionLocal() as db:
                pid.append(db.scalar(text("select pg_backend_pid()")))
                closing.set()
                return queue.close_session(db, db.get(QueueSession, session.id))

        with ThreadPoolExecutor(max_workers=2) as pool:
            joined = pool.submit(join)
            try:
                self.assertTrue(ready.wait(timeout=5))
                closed = pool.submit(close)
                self.assertTrue(closing.wait(timeout=5))
                # Проверяем настоящее ожидание замка PostgreSQL, а не скорость потока.
                deadline = time.monotonic() + 5
                blocked = False
                while time.monotonic() < deadline and not closed.done():
                    with engine.connect() as connection:
                        blocked = bool(connection.scalar(text("select pg_blocking_pids(:pid)"), {"pid": pid[0]}))
                    if blocked:
                        break
                    time.sleep(0.02)
                self.assertTrue(blocked, "Закрытие должно дождаться незавершённой записи")
            finally:
                release.set()
            self.assertIsInstance(joined.result(timeout=5), int)
            self.assertEqual(closed.result(timeout=5), 1)

    def test_parallel_joins_get_distinct_consecutive_numbers(self):
        session = self.open()
        barrier = threading.Barrier(6)

        def join(index):
            with SessionLocal() as db:
                current = db.get(QueueSession, session.id)
                barrier.wait(timeout=5)
                return self.add(db, current, f"Тестов Студент{index}").number

        with ThreadPoolExecutor(max_workers=6) as pool:
            numbers = list(pool.map(join, range(6)))
        self.assertEqual(sorted(numbers), list(range(1, 7)))

    def test_stale_and_unbound_forms_do_not_join_new_session(self):
        old = self.open()
        status, _, body = http("/join")
        self.assertEqual(status, 200)
        self.assertIn(f'name="session_id" value="{old.id}"', body.decode())
        challenge = self.captcha()
        self.close(old)
        current = self.open()
        data = self.form(old, challenge)
        status, _, body = http("/join", data)
        self.assertEqual(status, 409)
        self.assertIn("Приём изменился", body.decode())
        self.assertIn(f'name="session_id" value="{current.id}"', body.decode())
        self.assertIn('value="Тестов Иван"', body.decode())
        data.pop("session_id")
        self.assertEqual(http("/join", data)[0], 409)
        self.assertEqual(self.count(old), 0)
        self.assertEqual(self.count(current), 0)

    def test_student_cannot_leave_archived_session(self):
        session = self.open()
        with SessionLocal() as db:
            entry = self.add(db, db.get(QueueSession, session.id))
            cookie = "au_entry=" + write_entry_token(entry.token)
        self.close(session)
        self.assertEqual(http("/leave", {}, cookie)[0], 303)
        with SessionLocal() as db:
            actual = db.get(QueueEntry, entry.id)
            self.assertEqual(actual.status, EntryStatus.waiting)
            self.assertIsNone(actual.closed_at)

    def test_leave_does_not_overwrite_teacher_mark(self):
        session = self.open()
        with SessionLocal() as student, SessionLocal() as teacher:
            entry = self.add(student, student.get(QueueSession, session.id))
            queue.set_status(teacher, teacher.get(QueueEntry, entry.id), EntryStatus.done)
            self.assertEqual(entry.status, EntryStatus.waiting)
            self.assertFalse(queue.leave_queue(student, entry.token))
        with SessionLocal() as db:
            self.assertEqual(db.get(QueueEntry, entry.id).status, EntryStatus.done)

    def test_export_formula_cells_remain_text(self):
        values = ["=1+1", "+1+1", "-1+1", "@SUM(1)", " \t=1+1", "\ttext", "#N/A", 'текст; "пример"\nстрока']
        entries = [QueueEntry(number=i, full_name="Тестов Иван", group_name="КТ-24-04",
                              purpose=Purpose(title=value), comment=value,
                              status=EntryStatus.waiting, created_at=queue.utc_now())
                   for i, value in enumerate(values, start=1)]
        workbook = load_workbook(io.BytesIO(export.to_xlsx(entries, "=1+1")), data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet["A1"].data_type, "s")
        for row, value in enumerate(values, start=4):
            for column in (4, 5):
                cell = sheet.cell(row, column)
                self.assertEqual(cell.value, value)
                self.assertEqual(cell.data_type, "s")
        rows = list(csv.reader(io.StringIO(export.to_csv(entries, "=1+1").decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(rows[0][0], "'=1+1")
        for row, value in zip(rows[2:], values):
            expected = "'" + value if value in values[:6] else value
            self.assertEqual(row[3:5], [expected, expected])

    def test_control_characters_in_new_and_legacy_comments(self):
        session = self.open()
        data = self.form(session, self.captcha(), comment="текст\x01\x00\uffff\nстрока")
        self.assertEqual(http("/join", data)[0], 303)
        with SessionLocal() as db:
            entry = db.scalar(select(QueueEntry).where(QueueEntry.session_id == session.id))
            self.assertEqual(entry.comment, "текст\nстрока")
            # Такие данные могли сохраниться до исправления.
            entry.comment = "старый\x01\uffff\nтекст"
            db.commit()
        status, _, data = http(f"/admin/sessions/{session.id}/export.xlsx", cookie=self.admin)
        self.assertEqual(status, 200)
        sheet = load_workbook(io.BytesIO(data)).active
        self.assertEqual(sheet["E4"].value, "старый\nтекст")

    def test_invisible_comment_does_not_satisfy_required_purpose(self):
        session = self.open()
        data = self.form(session, self.captcha(), comment="\x01\uffff")
        with SessionLocal() as db:
            data["purpose_id"] = db.scalar(select(Purpose.id).where(Purpose.needs_comment.is_(True)))
        self.assertEqual(http("/join", data)[0], 422)
        self.assertEqual(self.count(session), 0)

    def test_finished_entry_offers_teacher_help_instead_of_rejoin(self):
        session = self.open()
        with SessionLocal() as db:
            entry = self.add(db, db.get(QueueSession, session.id))
            cookie = "au_entry=" + write_entry_token(entry.token)
            for status in (EntryStatus.done, EntryStatus.left, EntryStatus.no_show):
                with self.subTest(status=status):
                    queue.set_status(db, entry, status)
                    code, _, body = http("/", cookie=cookie)
                    self.assertEqual(code, 200)
                    page = body.decode()
                    self.assertNotIn('href="/join"', page)
                    self.assertNotIn("Встать в очередь снова", page)
                    self.assertIn("преподавател", page)

    def test_closed_page_keeps_polling_configuration(self):
        code, _, body = http("/")
        self.assertEqual(code, 200)
        page = body.decode()
        self.assertIn("Сейчас приёма нет", page)
        self.assertIn('data-session-id=""', page)
        self.assertIn('data-session-revision=""', page)
        self.assertIn('data-poll="', page)

    def test_public_revision_changes_with_session_metadata(self):
        session = self.open()
        initial = json.loads(http("/api/queue")[2])
        self.assertIn(initial["session_revision"], http("/")[2].decode())
        for field, value in (("room", "1300"), ("note", "Новое сообщение"), ("time_from", dt.time(12, 30))):
            with SessionLocal() as db:
                setattr(db.get(QueueSession, session.id), field, value)
                db.commit()
            updated = json.loads(http("/api/queue")[2])
            self.assertNotEqual(updated["session_revision"], initial["session_revision"])
            initial = updated
        self.close(session)
        new = self.open()
        updated = json.loads(http("/api/queue")[2])
        self.assertEqual(updated["session_id"], new.id)
        self.assertNotEqual(updated["session_revision"], initial["session_revision"])

    def test_custom_group_format_does_not_enable_standard_mask(self):
        custom = replace(settings, group_pattern=r"^[А-ЯЁ]{3}-[0-9]{2}-[0-9]{2}$", group_placeholder="КСП-24-04")
        self.assertEqual(custom.group_input_mask, "")
        self.assertEqual(settings.group_input_mask, "aa-00-00")
        session = self.open()
        challenge = self.captcha(solved=False)
        page = templates.get_template("join.html").render(
            settings=custom, session=session, challenge=challenge, errors={}, values={},
            purposes=[], groups=[], group_placeholder=custom.group_placeholder,
        )
        self.assertIn('data-group-input=""', page)
        self.assertIn("КСП-24-04", page)
        self.assertNotIn("две буквы", page)

    # ── Конец приёма по часам ──────────────────────────────────────────

    def test_session_end_follows_the_day_it_opened(self):
        session = self.open()
        with SessionLocal() as db:
            stored = db.get(QueueSession, session.id)
            opened = queue.now_local().replace(hour=22, minute=0, second=0, microsecond=0)
            stored.opened_at = opened.astimezone(dt.timezone.utc)

            stored.time_to = dt.time(23, 0)
            db.commit()
            ends_at = queue.session_ends_at(stored)
            self.assertEqual(ends_at, opened + dt.timedelta(hours=1))
            self.assertFalse(queue.joining_closed(stored, now=ends_at - dt.timedelta(minutes=1)))
            self.assertTrue(queue.joining_closed(stored, now=ends_at))

            # Приём через полночь: «до» раньше открытия — значит, конец завтра.
            stored.time_to = dt.time(0, 30)
            db.commit()
            self.assertEqual(
                queue.session_ends_at(stored),
                (opened + dt.timedelta(days=1)).replace(hour=0, minute=30),
            )

            # Без «до» приём по часам не заканчивается вовсе.
            stored.time_to = None
            db.commit()
            self.assertIsNone(queue.session_ends_at(stored))
            self.assertFalse(queue.joining_closed(stored))

    def test_expired_window_closes_joining_but_keeps_the_queue(self):
        session = self.open()
        with SessionLocal() as db:
            waiting = self.add(db, session)
            db.commit()
            number = waiting.number
        self.set_window(session)

        # Очередь жива: ждущий на месте, приём не ушёл в историю.
        with SessionLocal() as db:
            self.assertIsNotNone(queue.current_session(db))
        code, headers, _ = http("/join")
        self.assertEqual(code, 303)
        self.assertEqual(headers["Location"], "/")

        code, _, body = http("/")
        self.assertEqual(code, 200)
        self.assertIn("Запись закрыта", body.decode())
        self.assertIn(f">{number}<".encode(), body)

        # И прямой POST по открытой заранее форме тоже не проходит.
        code, headers, _ = http("/join", self.form(session, self.captcha(), "Опоздавший Пётр"))
        self.assertEqual(code, 303)
        self.assertEqual(headers["Location"], "/")
        self.assertEqual(self.count(session), 1)

    def test_session_closes_itself_once_the_last_person_is_served(self):
        session = self.open()
        with SessionLocal() as db:
            entry = self.add(db, session)
            db.commit()
            entry_id = entry.id
        self.set_window(session)

        with SessionLocal() as db:
            self.assertIsNotNone(queue.current_session(db), "с ждущим приём остаётся открытым")
            queue.set_status(db, db.get(QueueEntry, entry_id), EntryStatus.done)
        with SessionLocal() as db:
            self.assertIsNone(queue.current_session(db), "очередь опустела — приём закрылся сам")
            closed = db.get(QueueSession, session.id)
            self.assertIsNotNone(closed.closed_at)
            # Запись осталась в истории со своим статусом.
            self.assertEqual(db.get(QueueEntry, entry_id).status, EntryStatus.done)

    def test_open_session_without_end_time_never_closes_itself(self):
        session = self.open()
        with SessionLocal() as db:
            self.assertIsNone(db.get(QueueSession, session.id).time_to)
            self.assertIsNotNone(queue.current_session(db))
            self.assertIsNone(db.get(QueueSession, session.id).closed_at)

    def test_teacher_can_extend_the_window_instead_of_losing_the_session(self):
        session = self.open()
        with SessionLocal() as db:
            self.add(db, session)
            db.commit()
        self.set_window(session)

        later = (queue.now_local() + dt.timedelta(hours=1)).strftime("%H:%M")
        code, _, _ = http(
            "/admin/session/update",
            {"room": "1215", "time_from": "", "time_to": later, "note": ""},
            cookie=self.admin,
        )
        self.assertEqual(code, 303)
        with SessionLocal() as db:
            reopened = queue.current_session(db)
            self.assertIsNotNone(reopened)
            self.assertFalse(queue.joining_closed(reopened), "запись снова открыта")
        self.assertEqual(http("/join")[0], 200)

    def test_static_is_revalidated_after_a_rebuild(self):
        code, headers, _ = http("/static/js/app.js")
        self.assertEqual(code, 200)
        # Без этого браузер несколько часов рисует новую страницу старым
        # скриптом: разметка свежая, обработчиков новых заданий в ней нет.
        self.assertEqual(headers.get("Cache-Control"), "no-cache")
        self.assertTrue(headers.get("ETag"))

    def test_admin_captcha_has_one_shared_slot(self):
        code, _, body = http("/admin/captcha?kind=quiz", cookie=self.admin)
        self.assertEqual(code, 200)
        self.assertEqual(body.count(b"data-captcha-slot"), 1)
        self.assertIn(b"data-captcha-preview", body)
        self.assertIn(b"data-captcha-reload", body)


if __name__ == "__main__":
    unittest.main()
