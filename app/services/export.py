"""Выгрузка списка записанных: CSV для Excel и настоящий XLSX."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import Booking, BookingStatus
from app.services.slots import to_local

HEADERS = ["№", "Время", "ФИО", "Группа", "Цель визита", "Комментарий", "Статус", "Записан"]

BRAND = "015D1E"


def _rows(bookings: list[Booking]) -> list[list[str]]:
    rows: list[list[str]] = []
    for index, booking in enumerate(bookings, start=1):
        slot = booking.slot
        rows.append(
            [
                str(index),
                f"{to_local(slot.starts_at):%H:%M}–{to_local(slot.ends_at):%H:%M}",
                booking.full_name,
                booking.group_name,
                booking.purpose.title if booking.purpose else "",
                booking.comment,
                BookingStatus(booking.status).label,
                f"{to_local(booking.created_at):%d.%m.%Y %H:%M}",
            ]
        )
    return rows


def to_csv(bookings: list[Booking], title: str) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([title])
    writer.writerow(HEADERS)
    writer.writerows(_rows(bookings))
    # BOM — чтобы Excel не превратил кириллицу в кракозябры
    return "﻿".encode("utf-8") + buffer.getvalue().encode("utf-8")


def to_xlsx(bookings: list[Booking], title: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Запись на приём"

    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=13, color=BRAND)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    header_fill = PatternFill("solid", fgColor="E8F1EA")
    thin = Side(style="thin", color="D9DFDB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=3, column=column, value=header)
        cell.font = Font(bold=True, color="0B0F0C")
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(_rows(bookings), start=4):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=column in (5, 6))

    widths = [5, 14, 34, 12, 26, 30, 12, 18]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.freeze_panes = "A4"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
